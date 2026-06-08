"""
知识库处理模块（离线版本）
直接从本地 Excel 文件读取数据，无需网络
"""
import os
import openpyxl
from typing import List, Dict, Tuple, Optional
import logging

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# 本地文件路径
DATA_DIR = os.path.join(os.path.dirname(__file__), 'data')
ANSWER_FILE = os.path.join(DATA_DIR, 'answer.xlsx')
QUESTION_FILE = os.path.join(DATA_DIR, 'question.xlsx')


def load_questions_from_excel() -> List[str]:
    """从问题表加载所有问题"""
    questions = []
    try:
        wb = openpyxl.load_workbook(QUESTION_FILE)
        ws = wb.active
        for row in ws.iter_rows(min_row=2, values_only=True):
            if len(row) >= 2 and row[1] and isinstance(row[1], str):
                question = row[1].strip()
                if question:
                    questions.append(question)
        wb.close()
        logger.info(f"加载了 {len(questions)} 个问题")
        return questions
    except Exception as e:
        logger.error(f"读取问题表失败: {e}")
        return []


def load_answers_from_excel() -> List[Dict[str, str]]:
    """从答案表加载问题和答案"""
    qa_pairs = []
    try:
        wb = openpyxl.load_workbook(ANSWER_FILE)
        ws = wb.active
        for row in ws.iter_rows(min_row=2, values_only=True):
            if len(row) >= 2 and row[1] and isinstance(row[1], str):
                answer = row[1].strip()
                if answer:
                    qa_pairs.append({
                        'question_id': row[0],
                        'answer': answer
                    })
        wb.close()
        logger.info(f"加载了 {len(qa_pairs)} 个问答对")
        return qa_pairs
    except Exception as e:
        logger.error(f"读取答案表失败: {e}")
        return []


def find_answer_by_question(question: str) -> Tuple[Optional[str], Optional[str]]:
    """根据问题查找答案（模糊匹配）"""
    question = question.strip().lower()
    for qa in QA_PAIRS:
        stored_question = get_question_by_id(qa['question_id']).lower()
        if stored_question and (question in stored_question or stored_question in question):
            return qa['answer'], stored_question
    return None, None


def find_similar_questions(question: str, exclude: str = None, limit: int = 5) -> List[str]:
    """查找相似问题"""
    question = question.strip().lower()
    similar = []
    keywords = extract_keywords(question)
    for q in QUESTIONS:
        if exclude and (q.lower() in exclude.lower() or exclude.lower() in q.lower()):
            continue
        q_lower = q.lower()
        match_count = sum(1 for kw in keywords if kw in q_lower)
        score = match_count * 10
        if question.startswith(q_lower[:10]) if len(q_lower) >= 10 else False:
            score += 5
        if score > 0:
            similar.append((q, score))
    similar.sort(key=lambda x: x[1], reverse=True)
    return [q for q, _ in similar[:limit]]


def extract_keywords(text: str) -> List[str]:
    """提取文本关键词"""
    stop_words = {'的', '是', '了', '吗', '呢', '吧', '啊', '请问', '这个', '那个', '什么', '如何', '怎样', '有没有'}
    text = text.replace('？', ' ').replace('?', ' ').replace('。', ' ')
    words = text.split()
    return [w for w in words if len(w) >= 2 and w not in stop_words]


# 全局变量存储知识库数据
QUESTIONS: List[str] = []
QA_PAIRS: List[Dict[str, str]] = []


def get_question_by_id(question_id: str) -> Optional[str]:
    """根据问题ID获取问题内容"""
    try:
        idx = int(question_id.replace('问题', '')) - 1
        if 0 <= idx < len(QUESTIONS):
            return QUESTIONS[idx]
    except:
        pass
    return None


def init_knowledge_base():
    """初始化知识库"""
    global QUESTIONS, QA_PAIRS
    logger.info("正在初始化知识库...")
    QUESTIONS = load_questions_from_excel()
    QA_PAIRS = load_answers_from_excel()
    logger.info("知识库初始化完成")
    return len(QUESTIONS) > 0 and len(QA_PAIRS) > 0


def get_all_questions(page: int = 0, page_size: int = 10) -> Dict:
    """分页获取所有问题"""
    total = len(QUESTIONS)
    start = page * page_size
    end = start + page_size
    return {
        'success': True,
        'questions': QUESTIONS[start:end],
        'total': total,
        'page': page,
        'page_size': page_size,
        'total_pages': (total + page_size - 1) // page_size
    }


def get_paginated_questions(page: int = 0, page_size: int = 10) -> Dict:
    """获取分页问题（兼容旧API）"""
    return get_all_questions(page, page_size)


def search_answer(question: str) -> Dict:
    """搜索答案"""
    answer, matched_question = find_answer_by_question(question)
    if answer:
        return {
            'success': True,
            'answer': answer,
            'source': 'knowledge_base',
            'question': matched_question or question
        }
    return {
        'success': True,
        'answer': '抱歉，知识库中暂未收录该问题的答案。',
        'source': 'knowledge_base',
        'question': question
    }


def get_recommendations(question: str, limit: int = 5) -> List[str]:
    """获取推荐问题"""
    return find_similar_questions(question, exclude=question, limit=limit)


if __name__ == '__main__':
    if init_knowledge_base():
        print(f"知识库加载成功，共 {len(QUESTIONS)} 个问题")
        test_q = "用法用量"
        result = search_answer(test_q)
        print(f"搜索「{test_q}」: {result['answer']}")
        recs = get_recommendations(test_q)
        print(f"推荐问题: {recs}")
    else:
        print("知识库加载失败，请检查数据文件")