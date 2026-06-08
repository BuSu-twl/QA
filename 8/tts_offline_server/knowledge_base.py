"""
知识库处理模块
下载并处理答案表和问题表 Excel 文件
"""
import os
import requests
import openpyxl
from io import BytesIO
from typing import List, Dict, Tuple, Optional
import logging

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


# 本地文件路径
DATA_DIR = os.path.join(os.path.dirname(__file__), 'data')
ANSWER_FILE = os.path.join(DATA_DIR, 'answer.xlsx')
QUESTION_FILE = os.path.join(DATA_DIR, 'question.xlsx')


def ensure_data_dir():
    """确保数据目录存在"""
    if not os.path.exists(DATA_DIR):
        os.makedirs(DATA_DIR)


def download_file(url: str, local_path: str) -> bool:
    """下载文件到本地"""
    try:
        logger.info(f"正在下载文件: {url}")
        response = requests.get(url, timeout=30)
        response.raise_for_status()

        ensure_data_dir()
        with open(local_path, 'wb') as f:
            f.write(response.content)

        logger.info(f"文件已保存到: {local_path}")
        return True
    except Exception as e:
        logger.error(f"下载文件失败: {e}")
        return False


def load_questions_from_excel() -> List[str]:
    """从问题表加载所有问题"""
    questions = []

    # 如果本地文件不存在，先下载
    if not os.path.exists(QUESTION_FILE):
        download_file(QUESTION_URL, QUESTION_FILE)

    try:
        wb = openpyxl.load_workbook(QUESTION_FILE)
        # 获取第一个工作表
        ws = wb.active

        # 遍历所有行，获取第二列的问题内容（跳过表头行，从第2行开始）
        # 问题表结构：A列=问题序号，B列=问题内容
        for row in ws.iter_rows(min_row=2, values_only=True):
            if len(row) >= 2 and row[1] and isinstance(row[1], str):
                question = row[1].strip()
                if question:
                    questions.append(question)

        wb.close()
        logger.info(f"加载了 {len(questions)} 个问题")
        return questions
    except Exception as e:
        logger.error(f"读取问题表失败: {e}")
        return []


def load_answers_from_excel() -> List[Dict[str, str]]:
    """从答案表加载问题和答案"""
    qa_pairs = []

    # 如果本地文件不存在，先下载
    if not os.path.exists(ANSWER_FILE):
        download_file(ANSWER_URL, ANSWER_FILE)

    try:
        wb = openpyxl.load_workbook(ANSWER_FILE)
        ws = wb.active

        # 遍历所有行，假设第一列是问题序号答案（如"问题1答案"），第二列是答案
        # 答案表结构：A列=问题序号答案，B列=答案内容
        for row in ws.iter_rows(min_row=1, values_only=True):
            if len(row) >= 2 and row[0] and row[1]:
                # 提取问题序号（如从"问题1答案"提取"问题1"）
                question_key = str(row[0]).strip()  # 如"问题1答案"
                answer = str(row[1]).strip()
                if question_key and answer:
                    qa_pairs.append({
                        'question_key': question_key,  # 如"问题1答案"
                        'question_num': question_key.replace('答案', ''),  # 如"问题1"
                        'answer': answer
                    })

        wb.close()
        logger.info(f"加载了 {len(qa_pairs)} 个问答对")
        return qa_pairs
    except Exception as e:
        logger.error(f"读取答案表失败: {e}")
        return []


def find_best_match(user_question: str, qa_pairs: List[Dict[str, str]]) -> Optional[str]:
    """根据用户问题找到最佳匹配的答案"""
    if not qa_pairs:
        return None

    # 先从问题表加载问题列表进行匹配
    questions_data = []
    if os.path.exists(QUESTION_FILE):
        try:
            wb = openpyxl.load_workbook(QUESTION_FILE)
            ws = wb.active
            # 问题表结构：A列=问题序号，B列=问题内容
            for row in ws.iter_rows(min_row=2, values_only=True):
                if len(row) >= 2 and row[0] and row[1]:
                    question_num = str(row[0]).strip()  # 如"问题1"
                    question_content = str(row[1]).strip()  # 如"该处方是由哪些药材组成的?"
                    questions_data.append({
                        'question_num': question_num,
                        'question_content': question_content
                    })
            wb.close()
        except Exception as e:
            logger.error(f"读取问题表失败: {e}")

    user_q_lower = user_question.lower()
    best_answer = None
    best_score = 0

    # 遍历问题表中的每个问题，找到最佳匹配
    for q_data in questions_data:
        q_content = q_data['question_content']
        q_num = q_data['question_num']
        q_content_lower = q_content.lower()

        score = 0

        # 精确匹配
        if user_q_lower == q_content_lower:
            score = 100
        # 问题内容包含匹配
        elif q_content_lower in user_q_lower or user_q_lower in q_content_lower:
            score = 50
        else:
            # 分词匹配
            user_words = set(user_q_lower.replace('?', '').replace('？', '').replace('!', '！').split())
            q_words = set(q_content_lower.replace('?', '').replace('？', '').replace('!', '！').split())
            common_words = user_words & q_words
            score = len(common_words) * 2

        if score > best_score:
            best_score = score
            # 在答案表中查找对应的答案
            for qa in qa_pairs:
                if qa['question_num'] == q_num:
                    best_answer = qa['answer']
                    break

    # 如果没有匹配到，返回第一个答案作为默认（测试用）
    if not best_answer and qa_pairs:
        best_answer = qa_pairs[0]['answer']

    return best_answer


def get_paginated_questions(page: int = 0, page_size: int = 5) -> Tuple[List[str], int]:
    """获取分页的问题列表"""
    all_questions = load_questions_from_excel()
    total = len(all_questions)
    start = page * page_size
    end = start + page_size
    questions = all_questions[start:end]

    return questions, total


def get_recommended_questions(user_question: str, count: int = 5) -> List[str]:
    """根据用户输入推荐相关问题"""
    all_questions = load_questions_from_excel()
    if not all_questions:
        return []

    user_q_lower = user_question.lower().replace('?', '').replace('？', '').strip()
    scored_questions = []

    for q in all_questions:
        q_lower = q.lower().replace('?', '').replace('？', '').strip()

        # 排除与用户问题完全相同的问题
        if q_lower == user_q_lower:
            continue

        score = 0

        # 分词匹配
        user_words = set(user_q_lower.split())
        q_words = set(q_lower.split())

        # 共同词数量
        common_words = user_words & q_words
        score += len(common_words) * 3

        # 检查是否有共同的核心词（长度大于1的词）
        core_common_words = [w for w in common_words if len(w) > 1]
        if len(core_common_words) >= 1:
            score += 5

        # 检查是否有"该处方"这个开头（同一类问题）
        if user_q_lower.startswith('该处方') and q_lower.startswith('该处方'):
            score += 8

        # 关键词匹配
        keywords = ['如何', '怎么', '什么', '为什么', '哪里', '哪个', '多少', '时间', '方法', '步骤', '流程',
                   '适用', '组成', '用法', '用量', '注意', '事项', '常规', '处方', '服用', '饭前', '饭后']
        for kw in keywords:
            if kw in user_q_lower and kw in q_lower:
                score += 3

        if score > 0:
            scored_questions.append((q, score))

    # 按分数排序
    scored_questions.sort(key=lambda x: x[1], reverse=True)

    # 返回前 count 个
    return [q for q, _ in scored_questions[:count]]


def get_qa_pairs() -> List[Dict[str, str]]:
    """获取问答对列表（供API使用）"""
    return load_answers_from_excel()


def init_knowledge_base():
    """初始化知识库"""
    logger.info("正在初始化知识库...")
    load_questions_from_excel()
    load_answers_from_excel()
    logger.info("知识库初始化完成")
